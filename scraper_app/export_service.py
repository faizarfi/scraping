import io
import re
import logging
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from .category_normalizer import normalize_category

logger = logging.getLogger(__name__)

def resolve_district(place):
    """
    Pastikan kolom Kecamatan di Excel selalu terisi rapi:
    1. Ambil district tersimpan di database
    2. Fallback: Ekstrak nama kecamatan dari alamat lengkap via regex
    3. Fallback: Ekstrak dari query pencarian
    4. Default: '-' jika tidak ditemukan
    """
    if place.district and place.district.strip() and place.district.strip() != '-':
        return place.district.strip()

    combined = f"{place.address or ''} {place.search_query or ''}"

    # Deteksi regex umum: "Kecamatan Sukajadi", "Kec. Tebet", "Kecamatan Senen", dll.
    m = re.search(r'Kec(?:amatan|\.)\s+([A-Za-z0-9\s]+?)(?:,|$|\.|\d|\-)', combined, re.IGNORECASE)
    if m:
        name = m.group(1).strip()
        if 3 <= len(name) <= 30:
            return name.title()

    known_districts = [
        'Colomadu', 'Gondangrejo', 'Jaten', 'Jatipuro', 'Jatiyoso', 
        'Jenawi', 'Jumantono', 'Jumapolo', 'Karanganyar', 'Karangpandan', 
        'Kebakkramat', 'Kerjo', 'Matesih', 'Mojogedang', 'Ngargoyoso', 
        'Tasikmadu', 'Tawangmangu'
    ]
    for d in known_districts:
        if d.lower() in combined.lower():
            return d

    return "-"


OPERATIONAL_KEYWORDS = {
    'buka', 'tutup', 'open', 'closed', 'buka 24 jam', 'open 24 hours',
    'buka sekarang', 'open now', 'tutup sementara', 'temporarily closed',
    'tutup permanen', 'permanently closed', 'buka segera', 'akan segera buka',
    'dine-in', 'takeaway', 'delivery', 'makan di tempat', 'bawa pulang',
    'pesan antar', 'drive-through', 'antar tanpa bertemu', 'no-contact delivery'
}

NON_CATEGORY_TERMS = {
    'bersponsor', 'sponsored', 'iklan', 'ad', 'ads',
    'sesuai untuk keluarga', 'ramah anak', 'populer', 'terpopuler',
    'indonesia', 'jawa', 'jawa tengah', 'jawa barat', 'jawa timur',
    'bisnis', 'tempat',
    # Tombol aksi & label antarmuka Google Maps
    'rute', 'directions', 'direction', 'situs web', 'website',
    'simpan', 'save', 'saved', 'bagikan', 'share',
    'telepon', 'call', 'pesan', 'message', 'ringkasan', 'overview',
    'tentang', 'about', 'nearby', 'di sekitar', 'mulai', 'start',
    'foto', 'photo', 'photos', 'menu', 'ulasan', 'review', 'reviews'
}


def clean_text_glyphs(text):
    """
    Hapus karakter Private Use Area Unicode (ikon Google Maps: \\uE000-\\uF8FF),
    karakter kontrol, dan simbol bullet murni.
    """
    if not text:
        return ""
    cleaned = re.sub(r'[\ue000-\uf8ff]', '', text)
    cleaned = re.sub(r'[\x00-\x1f\x7f-\x9f]', '', cleaned)
    cleaned = re.sub(r'\s+', ' ', cleaned).strip()
    if re.match(r'^[\s\·\•\-\.\,\:\;\|\/\\]*$', cleaned):
        return ""
    return cleaned


def is_operational_or_status(text):
    if not text:
        return True
    t = clean_text_glyphs(text).strip().lower()
    if not t:
        return True
    if t in OPERATIONAL_KEYWORDS:
        return True
    if re.search(r'^(buka|tutup|open|closed)\b', t):
        return True
    if any(k in t for k in ['tutup pukul', 'buka pukul', 'buka 24 jam', 'closes at', 'opens at', 'closes soon', 'opens soon']):
        return True
    if re.search(r'\b(pukul|\bpm\b|\bam\b)\s*\d', t):
        return True
    return False


def is_rating_review_text(text):
    if not text:
        return False
    t = clean_text_glyphs(text).strip().lower()
    if not t:
        return True
    if any(phrase in t for phrase in [
        'tidak ada ulasan', 'belum ada ulasan', 'tanpa ulasan',
        'no reviews', 'no review', 'no rating'
    ]):
        return True
    if re.search(r'^\d+[\.\)]\s*(?:tidak ada ulasan|belum ada ulasan|no reviews?)', t):
        return True
    if re.search(r'^[1-5][,\.]\d\s*(?:\([0-9\.\,]+\))?$', t):
        return True
    if re.search(r'^\([0-9\.\,]+\)$', t):
        return True
    if re.search(r'^\d+[\.,]?\d*\s*(?:ulasan|reviews?|bintang|stars?)$', t):
        return True
    return False


def is_likely_address(text):
    if not text:
        return False
    t = clean_text_glyphs(text).strip()
    if not t or is_operational_or_status(t) or is_rating_review_text(t):
        return False
    t_lower = t.lower()
    if re.search(r'^hotel\s+bintang\s+\d', t_lower):
        return False
    if re.search(r'(?:rp\s*[\d\.]+|\/malam|per malam)', t_lower):
        return False
    if re.search(r'\b[2-9CFGHJMPQRVWX]{4,8}\+[2-9CFGHJMPQRVWX]{2,3}\b', t, re.I):
        return True
    if any(k in t_lower for k in [
        'jl.', 'jalan', 'raya', 'no.', 'rt.', 'rw.', 'gang', 'gg.',
        'kelurahan', 'kecamatan', 'kabupaten', 'kota', 'komplek',
        'blok', 'dusun', 'desa', 'perumahan', 'perum'
    ]):
        return True
    if re.search(r'\b\d{5}\b', t):
        return True
    if any(char.isdigit() for char in t) and len(t) >= 15 and not t_lower.startswith('hotel'):
        return True
    return False


def is_valid_category_candidate(text):
    if not text:
        return False
    t = clean_text_glyphs(text).strip()
    if not t or len(t) < 2 or len(t) > 50:
        return False
    t_lower = t.lower()
    if is_operational_or_status(t_lower) or is_rating_review_text(t_lower):
        return False
    if t_lower in NON_CATEGORY_TERMS:
        return False
    if (t.startswith('"') and t.endswith('"')) or (t.startswith("'") and t.endswith("'")) or '?' in t:
        return False
    if re.search(r'(?:rp\s*[\d\.]+|\/malam|per malam|\$\$\$?)', t_lower):
        return False
    if is_likely_address(t):
        return False
    if re.match(r'^[2-9A-Z]{4,8}\+?[2-9A-Z]{2,4}$', t, re.I):
        return False
    if len(t) >= 5 and re.match(r'^[A-Z0-9]+$', t) and not any(v in t_lower for v in ['a', 'i', 'u', 'e', 'o']):
        return False
    if not re.search(r'[a-zA-Z]', t):
        return False
    if re.search(r'^[\d\s\-\+\(\)]{8,}$', t):
        return False
    return True


def resolve_category(place):
    """
    Pastikan kolom Kategori di Excel selalu berisi nama bidang usaha/kategori yang benar dan terunifikasi,
    bukan teks 'Tidak ada ulasan', simbol bullet '·', tombol UI ('Rute'), Plus Code, atau status jam operasional.
    """
    search_q = (place.search_query or (place.job.query if place.job else "")).strip()

    candidate = ""
    cat = clean_text_glyphs(place.category or "").strip()
    if is_valid_category_candidate(cat):
        candidate = cat
    else:
        # Cek apakah kategori tersasar di kolom alamat (misal: 'Kantor Perusahaan · 93M3+Q7J' atau 'Hotel bintang 3')
        addr = clean_text_glyphs(place.address or "").strip()
        if '·' in addr or '•' in addr:
            parts = [p.strip() for p in re.split(r'[·•]', addr) if p.strip()]
            for p in parts:
                if is_valid_category_candidate(p):
                    candidate = p
                    break
        elif is_valid_category_candidate(addr) and not is_likely_address(addr):
            candidate = addr

        # Fallback ke kata kunci pencarian (search_query / job.query)
        if not candidate and search_q:
            clean_sq = re.split(r'\s+di\s+|\s+in\s+|,', search_q, flags=re.IGNORECASE)[0].strip()
            if is_valid_category_candidate(clean_sq):
                candidate = clean_sq

    if not candidate:
        candidate = "Bisnis"

    return normalize_category(candidate, search_q)


def resolve_address(place):
    """
    Pastikan alamat tidak memuat teks jam operasional, harga hotel, atau pecahan kategori
    """
    addr = clean_text_glyphs(place.address or "").strip()
    if not addr or is_operational_or_status(addr) or is_rating_review_text(addr):
        return "-"

    # Jika alamat memuat format 'Kategori · Alamat' (misal: 'Kantor Perusahaan · 93M3+Q7J')
    if '·' in addr or '•' in addr:
        parts = [p.strip() for p in re.split(r'[·•]', addr) if p.strip()]
        addr_parts = [p for p in parts if is_likely_address(p) or not is_valid_category_candidate(p)]
        if addr_parts:
            clean_res = ", ".join(addr_parts)
            if not is_operational_or_status(clean_res) and not is_rating_review_text(clean_res):
                return clean_res

    # Jika alamat adalah nama kategori hotel (seperti 'Hotel bintang 4') atau harga, kosongkan
    if re.search(r'^hotel\s+bintang\s+\d', addr.lower()) or re.search(r'(?:rp\s*[\d\.]+|\/malam)', addr.lower()):
        return "-"

    return addr


DEFAULT_EXPORT_COLUMNS = [
    'No', 'Nama Tempat', 'Kategori', 'Kecamatan', 'Rating', 'Jumlah Ulasan',
    'Alamat', 'No Telepon', 'Link WhatsApp', 'Website', 'Link Google Maps',
    'Latitude', 'Longitude', 'Kata Kunci', 'Waktu Ditemukan'
]


def build_dataframe_from_queryset(queryset):
    """
    Ubah queryset Place menjadi Pandas DataFrame yang rapi untuk diekspor.
    Menjamin 15 kolom terdefinisi meskipun data kosong (0 hasil).
    """
    data = []
    for idx, p in enumerate(queryset, start=1):
        wa_phone = p.clean_wa_phone
        wa_link = f"https://wa.me/{wa_phone}" if wa_phone else ""
        data.append({
            'No': idx,
            'Nama Tempat': (p.name or '').strip(),
            'Kategori': resolve_category(p),
            'Kecamatan': resolve_district(p),
            'Rating': p.rating if p.rating is not None else '',
            'Jumlah Ulasan': p.reviews_count if p.reviews_count is not None else 0,
            'Alamat': resolve_address(p),
            'No Telepon': (p.phone or '').strip(),
            'Link WhatsApp': wa_link,
            'Website': (p.website or '').strip(),
            'Link Google Maps': (p.google_maps_url or '').strip(),
            'Latitude': p.latitude if p.latitude is not None else '',
            'Longitude': p.longitude if p.longitude is not None else '',
            'Kata Kunci': (p.search_query or '').strip(),
            'Waktu Ditemukan': p.created_at.strftime('%Y-%m-%d %H:%M:%S') if p.created_at else ''
        })

    if not data:
        return pd.DataFrame(columns=DEFAULT_EXPORT_COLUMNS)

    df = pd.DataFrame(data)
    return df


def format_worksheet(worksheet, df):
    """
    Terapkan styling profesional pada worksheet: header tema gelap, border tipis,
    alignment angka/rating, auto-fit lebar kolom, dan link Google Maps yang dapat diklik.
    """
    if len(df.columns) == 0:
        return

    # Styling Header
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for col_num in range(1, len(df.columns) + 1):
        cell = worksheet.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Jika hanya ada header dan tidak ada baris data
    if df.empty:
        for col in worksheet.columns:
            col_letter = get_column_letter(col[0].column)
            worksheet.column_dimensions[col_letter].width = 18
        worksheet.row_dimensions[1].height = 28
        return

    # Baris isi data
    data_font = Font(name="Segoe UI", size=10)
    link_font = Font(name="Segoe UI", size=10, color="2563EB", underline="single")

    for row_idx, row in enumerate(worksheet.iter_rows(min_row=2, max_row=len(df) + 1, min_col=1, max_col=len(df.columns)), start=2):
        for col_idx, cell in enumerate(row, start=1):
            cell.font = data_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

            # Kolom Link Google Maps dan Link WhatsApp dibuat clickable hyperlink secara aman
            col_name = df.columns[col_idx - 1]
            if col_name in ['Link Google Maps', 'Link WhatsApp'] and cell.value and str(cell.value).startswith('http'):
                try:
                    cell.hyperlink = str(cell.value)
                    cell.font = link_font
                except Exception:
                    pass

            if col_name in ['Rating', 'Latitude', 'Longitude', 'No']:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_name == 'Jumlah Ulasan':
                cell.alignment = Alignment(horizontal="right", vertical="center")

    # Auto-adjust column widths
    for col in worksheet.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if len(val) > max_len:
                max_len = len(val)
        worksheet.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)

    worksheet.row_dimensions[1].height = 28


def generate_excel_bytes(queryset, title="Data Google Maps"):
    """
    Hasilkan file Excel (.xlsx) dengan styling rapi.
    Jika terdapat beberapa kategori berbeda, otomatis dibuatkan:
    - Tab Sheet 'Semua Data' (Master gabungan dengan 100% data)
    - Tab Sheet terpisah untuk masing-masing Kategori
    """
    df = build_dataframe_from_queryset(queryset)
    output = io.BytesIO()

    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        # Sheet 1: Master Sheet (Semua Data)
        master_sheet_name = 'Semua Data'
        df.to_excel(writer, index=False, sheet_name=master_sheet_name)
        
        master_ws = writer.sheets.get(master_sheet_name)
        if master_ws is not None:
            try:
                format_worksheet(master_ws, df)
            except Exception as e:
                logger.warning("Gagal format master sheet: %s", e)

        # Tab Sheet Per-Kategori (jika ada data dan kolom Kategori)
        if not df.empty and 'Kategori' in df.columns:
            try:
                # Normalisasi nama kategori (hapus spasi berlebih & Title Case)
                df_work = df.copy()
                df_work['Kategori_Norm'] = (
                    df_work['Kategori']
                    .fillna('Bisnis')
                    .astype(str)
                    .str.strip()
                    .str.title()
                )
                cat_counts = df_work['Kategori_Norm'].value_counts()
                
                # Jika kategori lebih dari 1, buat sub-sheet
                if len(cat_counts) > 1:
                    used_names_lower = {master_sheet_name.lower()}
                    
                    # Jika ada lebih dari 30 variasi kategori mikro, batasi 30 sheet utama
                    # dan sisanya ditampung dalam sheet 'Kategori Lainnya'
                    max_category_sheets = 30
                    top_cats = list(cat_counts.index[:max_category_sheets])
                    other_cats = list(cat_counts.index[max_category_sheets:])
                    
                    for cat_name in top_cats:
                        sub_df = df_work[df_work['Kategori_Norm'] == cat_name].drop(columns=['Kategori_Norm']).copy()
                        if sub_df.empty:
                            continue
                        
                        # Nomor urut baru per sheet kategori
                        sub_df['No'] = range(1, len(sub_df) + 1)
                        
                        # Bersihkan karakter terlarang Excel (\ / ? * : [ ]) dan tanda kutip
                        clean_name = re.sub(r'[\\/*?:\[\]]', '', str(cat_name)).strip().strip("'") or 'Kategori'
                        base_name = clean_name[:25]
                        sheet_name = base_name
                        counter = 2
                        while sheet_name.lower() in used_names_lower:
                            suffix = f"_{counter}"
                            sheet_name = f"{base_name[:31 - len(suffix)]}{suffix}"
                            counter += 1
                        used_names_lower.add(sheet_name.lower())
                        
                        sub_df.to_excel(writer, index=False, sheet_name=sheet_name)
                        
                        # Ambil worksheet secara aman (case-insensitive fallback)
                        ws = writer.sheets.get(sheet_name)
                        if ws is None and hasattr(writer, 'book'):
                            for title_str in writer.book.sheetnames:
                                if title_str.lower() == sheet_name.lower():
                                    ws = writer.book[title_str]
                                    break
                        if ws is not None:
                            try:
                                format_worksheet(ws, sub_df)
                            except Exception as err:
                                logger.warning("Gagal styling sheet %s: %s", sheet_name, err)
                                
                    # Jika terdapat kategori sisa (di atas 30)
                    if other_cats:
                        other_df = df_work[df_work['Kategori_Norm'].isin(other_cats)].drop(columns=['Kategori_Norm']).copy()
                        if not other_df.empty:
                            other_df['No'] = range(1, len(other_df) + 1)
                            clean_name = 'Kategori Lainnya'
                            sheet_name = clean_name
                            counter = 2
                            while sheet_name.lower() in used_names_lower:
                                sheet_name = f"{clean_name} {counter}"
                                counter += 1
                            used_names_lower.add(sheet_name.lower())
                            
                            other_df.to_excel(writer, index=False, sheet_name=sheet_name)
                            ws = writer.sheets.get(sheet_name)
                            if ws is None and hasattr(writer, 'book'):
                                for title_str in writer.book.sheetnames:
                                    if title_str.lower() == sheet_name.lower():
                                        ws = writer.book[title_str]
                                        break
                            if ws is not None:
                                try:
                                    format_worksheet(ws, other_df)
                                except Exception as err:
                                    logger.warning("Gagal styling sheet %s: %s", sheet_name, err)
            except Exception as e:
                logger.exception("Terjadi kesalahan saat memproses tab kategori Excel: %s", e)

    output.seek(0)
    return output.getvalue()



def generate_csv_bytes(queryset):
    """
    Hasilkan file CSV dalam format UTF-8 with BOM (utf-8-sig)
    sehingga langsung kompatibel di Microsoft Excel
    """
    df = build_dataframe_from_queryset(queryset)
    output = io.StringIO()
    # Gunakan utf-8-sig
    df.to_csv(output, index=False, encoding='utf-8-sig')
    return output.getvalue().encode('utf-8-sig')
