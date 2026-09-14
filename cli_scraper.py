"""
Google Maps Scraper CLI (Command Line Interface)
Jalankan scraping Google Maps secara langsung dari terminal tanpa perlu membuka web browser.
Contoh penggunaan:
    python cli_scraper.py --query "Cafe di Kemang Jakarta" --limit 30 --output cafe_kemang.xlsx
    python cli_scraper.py --query "Rumah Makan Padang Bandung" --limit 50 --output padang.csv
"""

import argparse
import re
import sys
import time
import urllib.parse
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from playwright.sync_api import sync_playwright

# Fix Windows console unicode issues
if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8', errors='replace')


def parse_rating_reviews(parent, lines):
    try:
        star_spans = parent.locator("span[role='img'], span[aria-label*='bintang'], span[aria-label*='star']").all()
        for s in star_spans:
            aria = s.get_attribute("aria-label") or ""
            match = re.search(r'([1-5][,\.]\d)\s*(?:bintang|stars?)(?:\s*([0-9\.\,]+))?', aria, re.I)
            if match:
                rating = float(match.group(1).replace(',', '.'))
                rev = int(re.sub(r'[^\d]', '', match.group(2))) if match.group(2) else 0
                return rating, rev
    except Exception:
        pass

    for line in lines:
        match = re.search(r'([1-5][,\.]\d)\s*\(([\d\.\,]+)\)', line)
        if match:
            try:
                rating = float(match.group(1).replace(',', '.'))
                reviews = int(re.sub(r'[^\d]', '', match.group(2))) if match.group(2) else 0
                return rating, reviews
            except Exception:
                pass
    return None, 0


def parse_coordinates(url):
    if not url:
        return None, None
    m1 = re.search(r'!3d(-?\d+\.\d+)!4d(-?\d+\.\d+)', url)
    if m1:
        return float(m1.group(1)), float(m1.group(2))
    m2 = re.search(r'@(-?\d+\.\d+),(-?\d+\.\d+)', url)
    if m2:
        return float(m2.group(1)), float(m2.group(2))
    return None, None

OPERATIONAL_KEYWORDS = {
    'buka', 'tutup', 'open', 'closed', 'buka 24 jam', 'open 24 hours',
    'buka sekarang', 'open now', 'tutup sementara', 'temporarily closed',
    'tutup permanen', 'permanently closed', 'buka segera', 'akan segera buka',
    'dine-in', 'takeaway', 'delivery', 'makan di tempat', 'bawa pulang',
    'pesan antar', 'drive-through', 'antar tanpa bertemu', 'no-contact delivery'
}

NON_CATEGORY_TERMS = {
    'bersponsor', 'sponsored', 'iklan', 'ad', 'ads',
    'sesuai untuk keluarga', 'ramah anak', 'populer', 'terpopuler',
    'indonesia', 'jawa', 'jawa tengah', 'jawa barat', 'jawa timur',
    'bisnis', 'tempat',
    # Tombol aksi & label antarmuka Google Maps
    'rute', 'directions', 'direction', 'situs web', 'website',
    'simpan', 'save', 'saved', 'bagikan', 'share',
    'telepon', 'call', 'pesan', 'message', 'ringkasan', 'overview',
    'tentang', 'about', 'nearby', 'di sekitar', 'mulai', 'start',
    'foto', 'photo', 'photos', 'menu', 'ulasan', 'review', 'reviews'
}

def clean_text_glyphs(text):
    if not text:
        return ""
    cleaned = re.sub(r'[\ue000-\uf8ff]', '', text)
    cleaned = re.sub(r'[\x00-\x1f\x7f-\x9f]', '', cleaned)
    cleaned = re.sub(r'\s+', ' ', cleaned).strip()
    if re.match(r'^[\s\·\•\-\.\,\:\;\|\/\\]*$', cleaned):
        return ""
    return cleaned

def is_operational_or_status(text):
    if not text:
        return True
    t = clean_text_glyphs(text).strip().lower()
    if not t or t in OPERATIONAL_KEYWORDS:
        return True
    if re.search(r'^(buka|tutup|open|closed)\b', t):
        return True
    if any(k in t for k in ['tutup pukul', 'buka pukul', 'buka 24 jam', 'closes at', 'opens at', 'closes soon', 'opens soon']):
        return True
    if re.search(r'\b(pukul|\bpm\b|\bam\b)\s*\d', t):
        return True
    return False

def is_rating_review_text(text):
    if not text:
        return False
    t = clean_text_glyphs(text).strip().lower()
    if not t:
        return True
    if any(phrase in t for phrase in [
        'tidak ada ulasan', 'belum ada ulasan', 'tanpa ulasan',
        'no reviews', 'no review', 'no rating'
    ]):
        return True
    if re.search(r'^\d+[\.\)]\s*(?:tidak ada ulasan|belum ada ulasan|no reviews?)', t):
        return True
    if re.search(r'^[1-5][,\.]\d\s*(?:\([0-9\.\,]+\))?$', t):
        return True
    if re.search(r'^\([0-9\.\,]+\)$', t):
        return True
    if re.search(r'^\d+[\.,]?\d*\s*(?:ulasan|reviews?|bintang|stars?)$', t):
        return True
    return False

def is_likely_address(text):
    if not text:
        return False
    t = clean_text_glyphs(text).strip()
    if not t or is_operational_or_status(t) or is_rating_review_text(t):
        return False
    t_lower = t.lower()
    if re.search(r'^hotel\s+bintang\s+\d', t_lower):
        return False
    if re.search(r'(?:rp\s*[\d\.]+|\/malam|per malam)', t_lower):
        return False
    if re.search(r'\b[2-9CFGHJMPQRVWX]{4,8}\+[2-9CFGHJMPQRVWX]{2,3}\b', t, re.I):
        return True
    if any(k in t_lower for k in [
        'jl.', 'jalan', 'raya', 'no.', 'rt.', 'rw.', 'gang', 'gg.',
        'kelurahan', 'kecamatan', 'kabupaten', 'kota', 'komplek',
        'blok', 'dusun', 'desa', 'perumahan', 'perum'
    ]):
        return True
    if re.search(r'\b\d{5}\b', t):
        return True
    if any(char.isdigit() for char in t) and len(t) >= 15 and not t_lower.startswith('hotel'):
        return True
    return False

def is_valid_category_candidate(text):
    if not text:
        return False
    t = clean_text_glyphs(text).strip()
    if not t or len(t) < 2 or len(t) > 50:
        return False
    t_lower = t.lower()
    if is_operational_or_status(t_lower) or is_rating_review_text(t_lower):
        return False
    if t_lower in NON_CATEGORY_TERMS:
        return False
    if (t.startswith('"') and t.endswith('"')) or (t.startswith("'") and t.endswith("'")) or '?' in t:
        return False
    if re.search(r'(?:rp\s*[\d\.]+|\/malam|per malam|\$\$\$?)', t_lower):
        return False
    if is_likely_address(t):
        return False
    if re.match(r'^[2-9A-Z]{4,8}\+?[2-9A-Z]{2,4}$', t, re.I):
        return False
    if len(t) >= 5 and re.match(r'^[A-Z0-9]+$', t) and not any(v in t_lower for v in ['a', 'i', 'u', 'e', 'o']):
        return False
    if not re.search(r'[a-zA-Z]', t):
        return False
    if re.search(r'^[\d\s\-\+\(\)]{8,}$', t):
        return False
    return True

def parse_category_and_address(lines, place_name=""):
    category = ""
    address = ""
    clean_lines = []
    p_name_lower = clean_text_glyphs(place_name or "").strip().lower()

    for line in lines:
        cleaned = clean_text_glyphs(line).strip()
        if not cleaned:
            continue
        if p_name_lower and cleaned.lower() == p_name_lower:
            continue
        if '·' not in cleaned and '•' not in cleaned:
            if is_operational_or_status(cleaned) or is_rating_review_text(cleaned):
                continue
        clean_lines.append(cleaned)

    for cleaned in clean_lines:
        if '·' in cleaned or '•' in cleaned:
            raw_parts = [clean_text_glyphs(p).strip() for p in re.split(r'[·•]', cleaned)]
            valid_parts = [p for p in raw_parts if p and not is_operational_or_status(p) and not is_rating_review_text(p)]
            for p in valid_parts:
                if not category and is_valid_category_candidate(p):
                    category = p
                    continue
                if not address and is_likely_address(p):
                    address = p
                    continue
        else:
            if not category and is_valid_category_candidate(cleaned):
                category = cleaned
            elif not address and is_likely_address(cleaned):
                address = cleaned

        if category and address:
            break

    if not address:
        for cleaned in clean_lines:
            if is_operational_or_status(cleaned) or is_rating_review_text(cleaned):
                continue
            if is_likely_address(cleaned):
                address = cleaned
                break

    if address and ('·' in address or '•' in address):
        p_addr = [clean_text_glyphs(x).strip() for x in re.split(r'[·•]', address) if clean_text_glyphs(x).strip()]
        for part in p_addr:
            if not category and is_valid_category_candidate(part):
                category = part
            elif is_likely_address(part):
                address = part

    if not category and address and is_valid_category_candidate(address) and not is_likely_address(address):
        category = address
        address = ""

    if not is_valid_category_candidate(category):
        category = ""

    if is_operational_or_status(address) or is_rating_review_text(address) or re.search(r'(?:rp\s*[\d\.]+|\/malam)', address.lower()):
        address = ""

    return category, address

def is_duplicate_cli(results, name, lat, lng, address):
    if not name:
        return True
    name_clean = name.strip().lower()

    if lat is not None and lng is not None:
        for item in results:
            if item.get('Nama Tempat', '').strip().lower() == name_clean:
                item_lat = item.get('Latitude')
                item_lng = item.get('Longitude')
                if isinstance(item_lat, (int, float)) and isinstance(item_lng, (int, float)):
                    if abs(item_lat - lat) < 0.00025 and abs(item_lng - lng) < 0.00025:
                        return True
        return False

    if address and address.strip():
        for item in results:
            if item.get('Nama Tempat', '').strip().lower() == name_clean:
                if item.get('Alamat', '').strip().lower() == address.strip().lower():
                    return True
        return False

    return False


def scrape_google_maps(query, district="", location="", limit=30, headless=True):
    # Susun kata kunci pencarian yang presisi
    parts = [query.strip()]
    if district and district.strip():
        d_clean = district.strip()
        if not d_clean.lower().startswith('kecamatan'):
            parts.append(f"Kecamatan {d_clean}")
        else:
            parts.append(d_clean)
    if location and location.strip():
        parts.append(location.strip())

    if len(parts) == 1:
        full_query = parts[0]
    elif len(parts) == 2:
        full_query = f"{parts[0]} di {parts[1]}"
    else:
        full_query = f"{parts[0]} di {parts[1]}, {parts[2]}"

    print(f"\n=======================================================")
    print(f"🗺️  Memulai scraping Google Maps: '{full_query}'")
    print(f"🎯 Target data: {limit} tempat")
    print(f"=======================================================\n")

    results = []
    seen_urls = set()

    with sync_playwright() as p:
        browser = p.chromium.launch(
            headless=headless,
            args=[
                "--disable-blink-features=AutomationControlled",
                "--no-sandbox",
                "--disable-setuid-sandbox"
            ]
        )
        context = browser.new_context(
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36",
            locale="id-ID",
            viewport={"width": 1280, "height": 850}
        )
        page = context.new_page()
        encoded = urllib.parse.quote_plus(full_query)
        search_url = f"https://www.google.com/maps/search/{encoded}/"
        
        print(f"[*] Membuka Google Maps...")
        page.goto(search_url, timeout=35000, wait_until="domcontentloaded")
        page.wait_for_timeout(3500)

        # Scroll loop
        max_scrolls = max(8, (limit // 5) + 5)
        print(f"[*] Melakukan scrolling untuk memuat daftar tempat...")

        for s in range(max_scrolls):
            page.mouse.move(250, 350)
            page.mouse.wheel(0, 3500)
            page.wait_for_timeout(1800)

            links = page.locator("div[role='feed'] a[href*='/maps/place/']").all()
            print(f"    -> Scroll #{s+1}: {len(links)} item terdeteksi...", end="\r")
            if len(links) >= limit:
                break

        print("\n[*] Mengekstrak rincian data tempat...")
        links = page.locator("div[role='feed'] a[href*='/maps/place/']").all()

        for idx, item in enumerate(links, start=1):
            if len(results) >= limit:
                break

            href = item.get_attribute("href") or ""
            if not href or href in seen_urls:
                continue
            seen_urls.add(href)

            name = item.get_attribute("aria-label") or ""
            if not name:
                name = item.inner_text().split("\n")[0].strip()
            if not name:
                continue

            parent = item.locator("xpath=..")
            parent_lines = [l.strip() for l in parent.inner_text().split("\n") if l.strip()]

            # Ekstraksi rating & ulasan
            rating, reviews = parse_rating_reviews(parent, parent_lines)

            category, address = parse_category_and_address(parent_lines)
            lat, lng = parse_coordinates(href)

            # Aturan Deduplikasi: HANYA skip jika Nama DAN Geolokasi SAMA
            if is_duplicate_cli(results, name, lat, lng, address):
                print(f"  [~] Dilewati (Duplikat Nama & Lokasi sama): {name}")
                continue

            results.append({
                'No': len(results) + 1,
                'Nama Tempat': name,
                'Kategori': category if category else query.strip(),
                'Kecamatan': district,
                'Rating': rating if rating is not None else '',
                'Jumlah Ulasan': reviews,
                'Alamat': address,
                'Link Google Maps': href,
                'Latitude': lat if lat is not None else '',
                'Longitude': lng if lng is not None else '',
                'Kata Kunci': full_query
            })
            print(f"  [+] #{len(results)}: {name} | ⭐ {rating or '-'} ({reviews}) | {category}")

        browser.close()

    print(f"\n[✓] Berhasil mengumpulkan {len(results)} tempat!")
    return results

def save_to_file(results, output_path):
    df = pd.DataFrame(results)
    if output_path.endswith('.csv'):
        df.to_csv(output_path, index=False, encoding='utf-8-sig')
        print(f"📄 Data berhasil disimpan ke file CSV: {output_path}")
    else:
        if not output_path.endswith('.xlsx'):
            output_path += '.xlsx'
        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            df.to_excel(writer, index=False, sheet_name='Tempat Google Maps')
            ws = writer.sheets['Tempat Google Maps']

            # Header styling
            fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
            font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
            for col in range(1, len(df.columns) + 1):
                c = ws.cell(row=1, column=col)
                c.fill = fill
                c.font = font
                c.alignment = Alignment(horizontal="center", vertical="center")

            # Column widths
            for col in ws.columns:
                max_len = max(len(str(c.value or '')) for c in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)

        print(f"📊 Data berhasil disimpan ke file Excel: {output_path}")

def main():
    parser = argparse.ArgumentParser(description="Google Maps Data Scraper CLI")
    parser.add_argument("-q", "--query", required=True, help="Kata kunci pencarian / Kategori (misal: 'Cafe', 'Restoran')")
    parser.add_argument("-k", "--kecamatan", "--district", default="", help="Nama Kecamatan (misal: 'Tebet', 'Kebayoran Baru')")
    parser.add_argument("-c", "--city", "--location", default="", help="Kota / Kabupaten (misal: 'Jakarta Selatan', 'Bandung')")
    parser.add_argument("-l", "--limit", type=int, default=30, help="Jumlah target tempat (default: 30)")
    parser.add_argument("-o", "--output", default="google_maps_data.xlsx", help="Nama file output (.xlsx atau .csv)")
    parser.add_argument("--headful", action="store_true", help="Buka browser dengan tampilan jendela (non-headless)")

    args = parser.parse_args()
    results = scrape_google_maps(
        query=args.query,
        district=args.kecamatan,
        location=args.city,
        limit=args.limit,
        headless=not args.headful
    )
    if results:
        save_to_file(results, args.output)
    else:
        print("⚠️ Tidak ada data tempat yang ditemukan.")

if __name__ == "__main__":
    main()

